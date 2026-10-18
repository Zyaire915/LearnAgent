import json
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\LearnAgent\presentation")
SOURCE = ROOT / "speech" / "speech.json"
OUTPUT = ROOT / "output" / "LearnAgent_答辩逐字演讲稿_v1.docx"
FONT = "Microsoft YaHei"
NAVY = "111A33"
PALE = "EEF3FB"
LINE = "D9D9D9"


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=LINE):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tbl_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


with SOURCE.open("r", encoding="utf-8") as handle:
    speeches = json.load(handle)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.45
normal.paragraph_format.space_after = Pt(6)

for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("LearnAgent 课程项目答辩逐字演讲稿")
set_font(title_run, 22, True, "000000")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("对应 11 页技术版答辩 PPT    建议总时长约 9 分 10 秒")
set_font(subtitle_run, 10.5, False, "666666")
subtitle.paragraph_format.space_after = Pt(16)

intro = doc.add_paragraph()
intro_run = intro.add_run("使用方法：")
set_font(intro_run, 11, True)
intro.add_run("正文可直接照读；“演示提示”不需要念出，用于提醒翻页、指图和停顿。正式答辩加入功能演示与成员分工后，可再统一压缩时长。")

doc.add_heading("时间安排", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(1.8), Cm(11.5), Cm(2.5)]
headers = ["页码", "页面主题", "建议时长"]
for idx, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
    cell.width = width
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, NAVY)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, 10.5, True, "FFFFFF")

for row_index, item in enumerate(speeches, start=1):
    cells = table.add_row().cells
    values = [str(item["slide"]), item["title"], item["duration"]]
    for column, (cell, value, width) in enumerate(zip(cells, values, widths)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            shade(cell, PALE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if column != 1 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(value)
        set_font(run, 10, False, "000000")
borders(table)
doc.add_page_break()

for index, item in enumerate(speeches):
    if index in (4, 8):
        doc.add_page_break()
    heading = doc.add_heading(f"第 {item['slide']} 页  {item['title']}", level=1)
    heading.paragraph_format.keep_with_next = True
    timing = doc.add_paragraph()
    timing.paragraph_format.keep_with_next = True
    run = timing.add_run(f"建议时长：{item['duration']}")
    set_font(run, 9.5, True, "3978E8")
    cue = doc.add_paragraph()
    cue.paragraph_format.keep_with_next = True
    lead = cue.add_run("演示提示：")
    set_font(lead, 10, True, "666666")
    cue_run = cue.add_run(item["cue"])
    set_font(cue_run, 10, False, "666666", True)
    script = doc.add_paragraph()
    script.paragraph_format.line_spacing = 1.55
    script.paragraph_format.space_after = Pt(12)
    script_run = script.add_run(item["script"])
    set_font(script_run, 11.5, False, "000000")

doc.add_heading("结束提醒", level=1)
closing = doc.add_paragraph()
closing.add_run("正式上台前至少完整计时练习两遍。第一遍保证内容准确，第二遍只保留自然停顿和指图动作。若总时长超过限制，优先压缩第 4 页和第 10 页，不要删除第 5 至第 9 页的 Java 实现说明。")

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_p.add_run("LearnAgent 课程项目答辩逐字稿")
set_font(footer_run, 8.5, False, "777777")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
